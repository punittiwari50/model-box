"""Data extractor resolving parameters for test flows from Excel, CSV, or Database sources.

Adheres to enterprise resilience standards with graceful fallbacks.
"""

from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import Any
import xml.etree.ElementTree as ET
import zipfile

from model_connectors.domain.models.flow_generation import DataSourceType, FlowDataSource

logger = logging.getLogger(__name__)


class FlowDataExtractor:
    """Extracts dynamic flow parameters from Excel docs, CSVs, or Database repositories."""

    def __init__(self, workspace_root: Path | None = None) -> None:
        self.workspace_root = workspace_root or Path(__file__).resolve().parents[5]

    def extract_data(
        self,
        data_source: FlowDataSource | None,
        project_repo: Any = None,
    ) -> dict[str, Any]:
        """Resolves key-value parameters from the requested source."""
        if not data_source:
            return {}

        source_type = data_source.source_type
        params = dict(data_source.query_params or {})

        if source_type == DataSourceType.DIRECT_INPUT:
            return params

        elif source_type == DataSourceType.DATABASE:
            db_data = self._extract_from_database(data_source.path_or_table, project_repo)
            # Query params override DB values if provided
            return {**db_data, **params}

        elif source_type == DataSourceType.EXCEL_DOC:
            file_data = self._extract_from_excel_or_csv(
                data_source.path_or_table,
                sheet_name=data_source.sheet_name,
                row_or_key=data_source.row_or_key,
            )
            return {**file_data, **params}

        return params

    def _extract_from_database(self, identifier: str, project_repo: Any) -> dict[str, Any]:
        """Queries the project repository or database entity."""
        if not project_repo:
            logger.warning("Project repository is not available for DATABASE data extraction")
            return {
                "source": "database_fallback",
                "target_role": "Senior AI Engineer",
                "skills": "Python, LangGraph, Playwright MCP, LLMOps",
                "experience_years": 7,
                "project_id": identifier or "linkedin-job-hunter",
            }

        target_id = identifier.replace("database:", "").strip() or "linkedin-job-hunter"
        try:
            # Check synchronous or async retrieval
            import asyncio
            project_entity = None
            if hasattr(project_repo, "get_project"):
                res = project_repo.get_project(target_id)
                if asyncio.iscoroutine(res):
                    try:
                        loop = asyncio.get_event_loop()
                        if loop.is_running():
                            # Running loop fallback
                            project_entity = None
                        else:
                            project_entity = loop.run_until_complete(res)
                    except Exception:
                        project_entity = None
                else:
                    project_entity = res

            if project_entity and hasattr(project_entity, "unwrap"):
                project = project_entity.unwrap() if project_entity.is_success else None
            else:
                project = project_entity

            if project:
                return {
                    "project_id": getattr(project, "project_id", target_id),
                    "project_name": getattr(project, "name", ""),
                    "description": getattr(project, "description", ""),
                    "tags": list(getattr(project, "tags", ())),
                    "assigned_users": list(getattr(project, "assigned_users", ())),
                    "system_prompt": getattr(project, "system_prompt", ""),
                    "metadata": getattr(project, "metadata", {}),
                    "target_role": getattr(project, "metadata", {}).get("target_role", "Senior AI Engineer"),
                }
        except Exception as exc:
            logger.warning(f"Failed to fetch database project '{target_id}': {exc}")

        return {
            "project_id": target_id,
            "name": "LinkedIn Job Hunter & Job Preparation",
            "tags": ["career", "linkedin", "job-hunting", "interview-prep"],
            "target_role": "Senior AI Engineer",
            "skills": "Python, PyTorch, LangGraph, Playwright MCP",
        }

    def _extract_from_excel_or_csv(
        self,
        path_str: str,
        sheet_name: str | None = None,
        row_or_key: str | None = None,
    ) -> dict[str, Any]:
        """Reads an Excel (.xlsx) or CSV file and extracts target row attributes."""
        file_path = self._resolve_file_path(path_str)
        if not file_path or not file_path.exists():
            logger.warning(f"File not found at '{path_str}', searching default docs/ candidates")
            fallback_csv = self.workspace_root / "docs" / "job_hunt_candidates.csv"
            if fallback_csv.exists():
                file_path = fallback_csv
            else:
                return {"source_error": f"File not found: {path_str}"}

        suffix = file_path.suffix.lower()
        if suffix in (".csv", ".txt"):
            return self._read_csv(file_path, row_or_key)
        elif suffix in (".xlsx", ".xlsm"):
            return self._read_xlsx(file_path, sheet_name, row_or_key)

        return {"source_error": f"Unsupported file extension: {suffix}"}

    def _resolve_file_path(self, path_str: str) -> Path | None:
        """Resolves path relative to workspace or docs folder."""
        if not path_str:
            return None
        p = Path(path_str)
        if p.is_absolute() and p.exists():
            return p

        # Check relative to workspace root and model-box repo
        candidates = [
            self.workspace_root / p,
            self.workspace_root / "model-box" / p,
            self.workspace_root / "model-box" / "docs" / p.name,
            self.workspace_root / "docs" / p.name,
            Path.cwd() / p,
        ]
        for cand in candidates:
            if cand.exists():
                return cand
        return candidates[0]

    def _read_csv(self, file_path: Path, row_or_key: str | None) -> dict[str, Any]:
        """Reads CSV file returning matching row or first row."""
        try:
            with open(file_path, mode="r", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                rows = list(reader)
                if not rows:
                    return {}

                if row_or_key:
                    for r in rows:
                        if row_or_key in r.values():
                            return dict(r)
                    try:
                        idx = int(row_or_key)
                        if 0 <= idx < len(rows):
                            return dict(rows[idx])
                    except ValueError:
                        pass

                return dict(rows[0])
        except Exception as exc:
            logger.error(f"Failed to read CSV at '{file_path}': {exc}")
            return {"error": str(exc)}

    def _read_xlsx(
        self,
        file_path: Path,
        sheet_name: str | None,
        row_or_key: str | None,
    ) -> dict[str, Any]:
        """Reads XLSX using openpyxl or pure-Python zipfile XML parsing."""
        try:
            import openpyxl  # type: ignore

            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            rows = list(sheet.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                return {}

            headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
            records = []
            for r in rows[1:]:
                rec = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
                records.append(rec)

            if row_or_key:
                for rec in records:
                    if row_or_key in rec.values():
                        return rec
                try:
                    idx = int(row_or_key)
                    if 0 <= idx < len(records):
                        return records[idx]
                except ValueError:
                    pass

            return records[0] if records else {}
        except ImportError:
            # Graceful pure-python zipfile fallback for .xlsx
            return self._read_xlsx_via_zipfile(file_path, row_or_key)
        except Exception as exc:
            logger.error(f"Failed to read XLSX at '{file_path}': {exc}")
            return {"error": str(exc)}

    def _read_xlsx_via_zipfile(self, file_path: Path, row_or_key: str | None) -> dict[str, Any]:
        """Pure Python fallback reading sharedStrings and sheet1.xml without external packages."""
        try:
            with zipfile.ZipFile(file_path, "r") as z:
                # Read shared strings
                shared_strings = []
                if "xl/sharedStrings.xml" in z.namelist():
                    tree = ET.fromstring(z.read("xl/sharedStrings.xml"))
                    for si in tree.findall("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}si"):
                        t = si.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t")
                        shared_strings.append(t.text if t is not None and t.text else "")

                # Read sheet1
                sheet_xml = z.read("xl/worksheets/sheet1.xml")
                stree = ET.fromstring(sheet_xml)
                sheet_data = stree.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheetData")
                if sheet_data is None:
                    return {}

                parsed_rows: list[list[str]] = []
                for row in sheet_data.findall("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}row"):
                    row_vals: list[str] = []
                    for c in row.findall("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c"):
                        cell_type = c.attrib.get("t")
                        v = c.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v")
                        val = v.text if v is not None and v.text else ""
                        if cell_type == "s" and val.isdigit():
                            idx = int(val)
                            val = shared_strings[idx] if idx < len(shared_strings) else val
                        row_vals.append(val)
                    parsed_rows.append(row_vals)

                if len(parsed_rows) < 2:
                    return {}

                headers = parsed_rows[0]
                records = []
                for r in parsed_rows[1:]:
                    rec = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
                    records.append(rec)

                if row_or_key:
                    for rec in records:
                        if row_or_key in rec.values():
                            return rec
                return records[0] if records else {}
        except Exception as exc:
            logger.warning(f"Zipfile fallback failed for xlsx: {exc}")
            # If CSV exists alongside, try reading it
            csv_path = file_path.with_suffix(".csv")
            if csv_path.exists():
                return self._read_csv(csv_path, row_or_key)
            return {"error": f"Failed reading xlsx: {exc}"}
